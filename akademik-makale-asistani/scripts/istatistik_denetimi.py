#!/usr/bin/env python3
"""Metindeki istatistiklerin iç tutarlılığını denetler (statcheck ve GRIM mantığı).

1. APA biçimli test sonuçlarından p değerini yeniden hesaplar: t(sd)=, F(sd1, sd2)=, r(sd)=,
   χ2(sd[, N=])=, z=. Bildirilen p ile uyumsuzluğu ve karar değişikliğini (α=.05) işaretler.
2. GRIM: tam sayı puanlı (Likert, madde toplamı) değişkende bildirilen ortalamanın n ile
   matematiksel olarak mümkün olup olmadığını sınar.
3. Yüzde–frekans tutarlılığı.

Kullanım:
    python3 istatistik_denetimi.py metin.txt|metin.docx|-
    import istatistik_denetimi as sd; sd.grim(3.47, 28, madde=1); sd.yuzde(17, 52, 32.7)
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats

SAYI = r"[-−–]?\d*[.,]?\d+"
P = r"p\s*([<>=≤≥])\s*(\.\d+|0[.,]\d+|1)"
DESEN = {
    "t": re.compile(rf"\bt\s*\(\s*(\d+(?:[.,]\d+)?)\s*\)\s*=\s*({SAYI})\s*,\s*{P}", re.I),
    "F": re.compile(rf"\bF\s*\(\s*(\d+(?:[.,]\d+)?)\s*,\s*(\d+(?:[.,]\d+)?)\s*\)\s*=\s*({SAYI})\s*,\s*{P}"),
    "r": re.compile(rf"\br\s*\(\s*(\d+)\s*\)\s*=\s*({SAYI})\s*,\s*{P}"),
    "chi2": re.compile(rf"(?:χ2|χ²|chi2|chi-square|X2)\s*\(\s*(\d+)(?:\s*,\s*N\s*=\s*\d+)?\s*\)\s*=\s*({SAYI})\s*,\s*{P}", re.I),
    "z": re.compile(rf"\bz\s*=\s*({SAYI})\s*,\s*{P}", re.I),
}


def _f(s):
    return float(str(s).replace("−", "-").replace("–", "-").replace(",", "."))


def _ondalik(s):
    s = str(s).replace(",", ".")
    return len(s.split(".")[1]) if "." in s else 0


def _p_hesap(tur, g):
    if tur == "t":
        df, t = _f(g[0]), _f(g[1]); return 2 * stats.t.sf(abs(t), df), g[2], g[3], g[1]
    if tur == "F":
        d1, d2, F = _f(g[0]), _f(g[1]), _f(g[2]); return stats.f.sf(F, d1, d2), g[3], g[4], g[2]
    if tur == "r":
        df, r = _f(g[0]), _f(g[1])
        if abs(r) >= 1:
            return np.nan, g[2], g[3], g[1]
        t = r * np.sqrt(df / (1 - r * r)); return 2 * stats.t.sf(abs(t), df), g[2], g[3], g[1]
    if tur == "chi2":
        df, x = _f(g[0]), _f(g[1]); return stats.chi2.sf(x, df), g[2], g[3], g[1]
    z = _f(g[0]); return 2 * stats.norm.sf(abs(z)), g[1], g[2], g[0]


def _uyumlu(hesap, op, rap, istatistik_str, tur, g):
    """İstatistiğin yuvarlama aralığını hesaba katarak bildirilen p ile uyumu sınar."""
    d = _ondalik(istatistik_str)
    yarim = 0.5 * 10 ** (-d) if d else 0.5
    deger = _f(istatistik_str)
    uclar = []
    for delta in (-yarim, yarim):
        gg = list(g)
        idx = {"t": 1, "F": 2, "r": 1, "chi2": 1, "z": 0}[tur]
        gg[idx] = str(deger + delta)
        try:
            uclar.append(_p_hesap(tur, gg)[0])
        except Exception:
            pass
    lo, hi = (min(uclar + [hesap]), max(uclar + [hesap])) if uclar else (hesap, hesap)
    rp = _f(rap)
    if op in ("<", "≤"):
        return lo < rp or np.isclose(lo, rp)
    if op in (">", "≥"):
        return hi > rp
    dp = _ondalik(rap)
    tol = 0.5 * 10 ** (-dp)
    return (lo - tol) <= rp <= (hi + tol)


def denetle(metin: str, alfa=0.05) -> pd.DataFrame:
    rows = []
    for tur, rx in DESEN.items():
        for m in rx.finditer(metin):
            g = m.groups()
            try:
                hesap, op, rap, ist = _p_hesap(tur, g)
            except Exception:
                continue
            if hesap != hesap:
                rows.append({"İfade": m.group(0), "Hesaplanan p": "—", "Durum": "r ≥ 1: geçersiz"}); continue
            ok = _uyumlu(hesap, op, rap, ist, tur, g)
            rp = _f(rap)
            rap_anl = (op in ("<", "≤") and rp <= alfa) or (op == "=" and rp < alfa)
            karar_degisti = (not ok) and (rap_anl != (hesap < alfa))
            rows.append({"İfade": m.group(0)[:70], "Hesaplanan p": f"{hesap:.4f}",
                         "Durum": "tutarlı" if ok else ("BÜYÜK HATA: anlamlılık kararı değişiyor" if karar_degisti else "tutarsız")})
    return pd.DataFrame(rows)


def grim(ortalama: float, n: int, madde: int = 1, ondalik: int | None = None) -> dict:
    """Tam sayı puanlı veride (madde sayısı = madde) ortalama ile n uyumlu mu?
    Yalnızca tam sayılardan oluşan ölçümler için geçerlidir; n çok büyükse (n*madde ≥ 10^ondalık) test bilgisiz kalır."""
    d = ondalik if ondalik is not None else _ondalik(str(ortalama))
    N = n * madde
    if N >= 10 ** d:
        return {"sonuç": "bilgisiz (n × madde ≥ 10^ondalık)", "tutarlı": None}
    yakin = round(ortalama * N)
    adaylar = [(k / N) for k in (yakin - 1, yakin, yakin + 1)]
    ok = any(round(a, d) == round(ortalama, d) for a in adaylar)
    return {"ortalama": ortalama, "n": n, "madde": madde, "tutarlı": ok,
            "olası ortalamalar": [round(a, d) for a in adaylar]}


def yuzde(frekans: int, n: int, bildirilen_yuzde: float, ondalik: int = 1) -> dict:
    gercek = frekans / n * 100
    return {"hesaplanan %": round(gercek, ondalik + 1), "bildirilen %": bildirilen_yuzde,
            "tutarlı": abs(round(gercek, ondalik) - bildirilen_yuzde) <= 10 ** (-ondalik) / 2 + 1e-9}


def _oku(yol):
    if yol == "-":
        return sys.stdin.read()
    p = Path(yol)
    if p.suffix.lower() == ".docx":
        import docx
        d = docx.Document(str(p))
        parcalar = [x.text for x in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parcalar.append(" | ".join(c.text for c in r.cells))
        return "\n".join(parcalar)
    return p.read_text(encoding="utf-8", errors="ignore")


if __name__ == "__main__":
    metin = _oku(sys.argv[1] if len(sys.argv) > 1 else "-")
    r = denetle(metin)
    if r.empty:
        print("APA biçiminde test istatistiği bulunamadı (ör. t(118) = 2.41, p = .018).")
    else:
        print(r.to_string(index=False))
        print(f"\nToplam {len(r)} ifade; tutarsız {int((r['Durum'] != 'tutarlı').sum())}.")
